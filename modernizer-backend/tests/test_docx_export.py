"""Word (.docx) downloads: converted from the same Markdown the .md downloads
return — headings, lists, tables, code, links and the UI Screens' images."""
import asyncio
import io
import json
import struct
import time
import zlib

from docx import Document
from fastapi.testclient import TestClient

import main
from agents.shared import docx_export as dx
from tests.test_stack_discovery_workflow import PRESCAN, _session


def _png(width: int, height: int) -> bytes:
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    raw = b"".join(b"\x00" + b"\xff\xff\xff" * width for _ in range(height))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def _doc(data: bytes):
    return Document(io.BytesIO(data))


MD = """## Executive Summary

The system **takes orders** and checks them in `OrderService` ([docs](https://example.com), [file](file:///etc/passwd)).

- one
  - two **bold**
    - three
1. first
3. third

| Rule | Source |
|---|---|
| over 100 \\| under 5 | `Order.java:12` |

> quoted

```text
  +--box--+
  |  a    |
```

<!-- SECTION: END -->

### Detail
"""


def test_markdown_structure_becomes_word_structure():
    doc = _doc(dx.to_docx("Reverse Engineering", [("Business Requirements", MD)]))
    styled = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
    assert styled[0] == ("Title", "Reverse Engineering")
    assert ("Heading 1", "Business Requirements") in styled
    assert ("Heading 2", "Executive Summary") in styled and ("Heading 3", "Detail") in styled
    assert ("List Bullet", "one") in styled and ("List Bullet 2", "two bold") in styled
    assert ("List Bullet 3", "three") in styled
    assert ("Normal", "1. first") in styled and ("Normal", "3. third") in styled      # the document's own numbers
    assert ("Quote", "quoted") in styled
    assert ("Normal", "  +--box--+\n  |  a    |") in styled                          # diagrams keep their spacing
    assert not any("SECTION" in text for _, text in styled)
    summary = next(p for p in doc.paragraphs if p.text.startswith("The system"))
    assert any(r.bold and r.text == "takes orders" for r in summary.runs)
    assert any(r.font.name == dx.MONO and r.text == "OrderService" for r in summary.runs)
    table = doc.tables[0]
    assert [[c.text for c in row.cells] for row in table.rows] == [["Rule", "Source"],
                                                                   ["over 100 | under 5", "Order.java:12"]]
    assert all(r.bold for r in table.rows[0].cells[0].paragraphs[0].runs)
    links = [r.target_ref for r in doc.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert links == ["https://example.com"]                                       # file:// is not linked


def test_images_are_embedded_only_from_the_files_given(tmp_path):
    shot = tmp_path / "SCR-0001.png"
    shot.write_bytes(_png(1366, 900))
    tall = tmp_path / "SCR-0002.png"
    tall.write_bytes(_png(1366, 4000))
    md = ("![SCR-0001 — default](screens/SCR-0001.png)\n\n![SCR-0002 — long](screens/SCR-0002.png)\n\n"
          "![secret](../../etc/passwd)\n\n![remote](https://example.com/x.png)\n")
    doc = _doc(dx.to_docx("UI Screens", [(None, md)], {"screens/SCR-0001.png": shot,
                                                        "screens/SCR-0002.png": tall}))
    shapes = doc.inline_shapes
    assert len(shapes) == 2
    assert round(shapes[0].width.inches, 2) == dx.PAGE_WIDTH_IN
    assert shapes[1].height.inches <= dx.MAX_IMAGE_HEIGHT_IN + 0.01                 # a long page fits the page
    texts = [p.text for p in doc.paragraphs]
    assert "[image: secret]" in texts and "[image: remote]" in texts


def test_control_characters_never_break_the_file():
    doc = _doc(dx.to_docx("T\x07", [(None, "bad \x00\x1b char `x\x01`\n\n| a\x02 |\n|---|\n| b\x03 |")]))
    # CommonMark itself turns NUL into U+FFFD; the other control characters are dropped.
    assert "bad \ufffd char x" in [p.text for p in doc.paragraphs]
    assert [[c.text for c in r.cells] for r in doc.tables[0].rows] == [["a"], ["b"]]


def test_a_large_rules_catalog_converts_quickly():
    rows = "\n".join(f"| BR-{i:04d} | Rule {i} **x** | `A.java:{i}` | a | b | c | d |" for i in range(2000))
    started = time.monotonic()
    data = dx.to_docx("Big", [(None, "| a | b | c | d | e | f | g |\n|---|---|---|---|---|---|---|\n" + rows)])
    assert time.monotonic() - started < 20
    assert len(_doc(data).tables[0].rows) == 2001


# ── endpoints ───────────────────────────────────────────────────────────────

def _session_with_documents(tmp_path, screens: bool = False) -> str:
    sid = _session(PRESCAN)
    delta = {"brd": "## Executive Summary\nOrders are checked.", "technical_spec": "## Architecture Overview\nOne WAR.",
             "test_inventory": "- `OrderTest.java` — JUnit test"}
    if screens:
        (tmp_path / "SCR-0001.png").write_bytes(_png(40, 30))
        delta.update({"ui_screens_dir": str(tmp_path),
                      "ui_screens": "## UI Screens\n\n#### SCR-0001 · default\n\n![SCR-0001 — default](screens/SCR-0001.png)\n",
                      "ui_screens_json": json.dumps({"screens": [{"id": "SCR-0001", "image": "SCR-0001.png"}]})})
    asyncio.run(main._update_state(sid, delta))
    return sid


def test_brd_downloads_as_markdown_by_default_and_as_word(tmp_path):
    sid = _session_with_documents(tmp_path)
    with TestClient(main.app) as client:
        md = client.get(f"/api/sessions/{sid}/download/brd")
        word = client.get(f"/api/sessions/{sid}/download/brd?format=docx")
        bad = client.get(f"/api/sessions/{sid}/download/brd?format=pdf")
    assert md.status_code == 200 and md.text.startswith("## Executive Summary")
    assert md.headers["content-disposition"].endswith('.md"')
    assert word.status_code == 200 and word.headers["content-type"] == dx.MEDIA_TYPE
    assert word.headers["content-disposition"].endswith('.docx"')
    assert "Orders are checked." in [p.text for p in _doc(word.content).paragraphs]
    assert bad.status_code == 400


def test_technical_documentation_downloads_in_both_formats(tmp_path):
    sid = _session_with_documents(tmp_path)
    with TestClient(main.app) as client:
        md = client.get(f"/api/sessions/{sid}/download/technical-spec?format=md")
        word = client.get(f"/api/sessions/{sid}/download/technical-spec?format=docx")
    assert "# Technical Specification" in md.text and "# Existing Test Inventory" in md.text
    assert "OrderTest.java" in md.text
    headings = [p.text for p in _doc(word.content).paragraphs if p.style.name == "Heading 1"]
    assert headings == ["Technical Specification", "Existing Test Inventory"]


def test_technical_documentation_404s_before_it_exists():
    sid = _session(PRESCAN)
    with TestClient(main.app) as client:
        assert client.get(f"/api/sessions/{sid}/download/technical-spec").status_code == 404


def test_whole_document_in_word_carries_the_screens(tmp_path):
    sid = _session_with_documents(tmp_path, screens=True)
    with TestClient(main.app) as client:
        md = client.get(f"/api/sessions/{sid}/download/reverse-engineering")
        word = client.get(f"/api/sessions/{sid}/download/reverse-engineering?format=docx")
    assert md.text.startswith("# Reverse Engineering") and "UI Screens" not in md.text     # the .md is unchanged
    doc = _doc(word.content)
    assert [p.text for p in doc.paragraphs if p.style.name == "Heading 1"] == [
        "Business Requirements", "Technical Specification", "Existing Test Inventory", "UI Screens"]
    assert len(doc.inline_shapes) == 1


def test_ui_screens_download_as_word_or_zip(tmp_path):
    sid = _session_with_documents(tmp_path, screens=True)
    with TestClient(main.app) as client:
        default = client.get(f"/api/sessions/{sid}/download/ui-screens")
        word = client.get(f"/api/sessions/{sid}/download/ui-screens?format=docx")
    assert default.headers["content-type"] == "application/zip"                            # unchanged default
    assert word.headers["content-type"] == dx.MEDIA_TYPE and len(_doc(word.content).inline_shapes) == 1


def test_property_elements_follow_the_schema_order():
    """Word refuses out-of-order properties (ECMA-376), even where other readers do not."""
    from docx.oxml.ns import qn
    doc = _doc(dx.to_docx("T", [(None, "```text\n  x\n```\n\n---\n\n| a |\n|---|\n| b |\n")]))
    for ppr in doc.element.body.iter(qn("w:pPr")):
        names = [c.tag.split("}")[1] for c in ppr]
        positions = [dx._PPR_ORDER.index(n) for n in names if n in dx._PPR_ORDER]
        assert positions == sorted(positions), names
    for tcpr in doc.element.body.iter(qn("w:tcPr")):
        names = [c.tag.split("}")[1] for c in tcpr]
        assert [dx._TCPR_ORDER.index(n) for n in names] == sorted(dx._TCPR_ORDER.index(n) for n in names), names


def test_a_section_does_not_repeat_its_own_title():
    doc = _doc(dx.to_docx("Doc", [("UI Screens", "## UI Screens\n\nintro\n\n### `GET /jobs`\n")]))
    heads = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert heads == [("Heading 1", "UI Screens"), ("Heading 2", "GET /jobs")]

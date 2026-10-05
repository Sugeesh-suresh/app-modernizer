"""
Word (.docx) versions of the reverse-engineering documents, converted from
their Markdown — no model involved: the Word file says exactly what the
Markdown says.

The Markdown is parsed with markdown-it (CommonMark plus GitHub tables and
strikethrough; raw HTML is not interpreted) and written with python-docx:

- headings → Word headings; paragraphs with **bold**, *italic*, `code`,
  ~~strikethrough~~ and links (http, https and mailto only become hyperlinks);
- bulleted lists → Word's bullet styles (three levels), numbered lists keep
  their own numbers;
- tables → bordered tables with a shaded header row;
- fenced code and the plain-text diagrams → monospace blocks, whitespace kept;
- block quotes, horizontal rules;
- images → embedded only when the caller supplies the file for that exact
  reference (the UI Screens' `screens/SCR-0001.png`); any other image becomes
  its alt text, so a document can never pull in a file it was not given.

Section markers (`<!-- SECTION: … -->`) and other HTML comments are dropped.
"""
import io
import re
import struct
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.shared import Inches, Pt, RGBColor
from markdown_it import MarkdownIt

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_MD = MarkdownIt("commonmark", {"html": False}).enable(["table", "strikethrough"])
_COMMENT_LINE = re.compile(r"^\s*<!--.*?-->\s*$", re.M)
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")   # not allowed in Word XML
_LINKABLE = re.compile(r"^(?:https?://|mailto:)", re.I)
MONO = "Consolas"
PAGE_WIDTH_IN = 7.0          # inside 0.75 in margins on US Letter, portrait
MAX_IMAGE_HEIGHT_IN = 8.6
LANDSCAPE_WIDTH_IN = 9.5     # the same margins, landscape
LANDSCAPE_IMAGE_HEIGHT_IN = 6.6
WIDE_TABLE_COLUMNS = 6       # a document with a table this wide is laid out in landscape


def _clean(text: str) -> str:
    return _CONTROL.sub("", text or "")


def _png_dpi(path: Path) -> float | None:
    """The horizontal resolution a PNG declares (its pHYs chunk), in dots per inch."""
    with open(path, "rb") as f:
        data = f.read(4096)
    at = data.find(b"pHYs")
    if at < 0 or len(data) < at + 13:
        return None
    x, _, unit = struct.unpack(">IIB", data[at + 4:at + 13])
    return x * 0.0254 if unit == 1 and x else None


def _png_size(path: Path) -> tuple[int, int] | None:
    with open(path, "rb") as f:
        head = f.read(24)
    if head[:8] == b"\x89PNG\r\n\x1a\n" and len(head) == 24:
        return struct.unpack(">II", head[16:24])
    return None


# Word rejects a file whose property elements are out of the schema's order
# (ECMA-376 CT_PPr / CT_TcPr), so new ones go before the first that must follow.
_PPR_ORDER = ("pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
              "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
              "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
              "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
              "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
              "sectPr", "pPrChange")
_TCPR_ORDER = ("cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
               "textDirection", "tcFitText", "vAlign", "hideMark", "headers", "cellIns", "cellDel", "cellMerge",
               "tcPrChange")


def _insert_ordered(props, child, order: tuple) -> None:
    name = child.tag.split("}")[1]
    later = {qn(f"w:{n}") for n in order[order.index(name) + 1:]}
    for existing in props:
        if existing.tag in later:
            existing.addprevious(child)
            return
    props.append(child)


def _shade(element, fill: str) -> None:
    cell = element.tag == qn("w:tc")
    props = element.get_or_add_tcPr() if cell else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    _insert_ordered(props, shd, _TCPR_ORDER if cell else _PPR_ORDER)


def _rule(paragraph) -> None:
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "A0A8B8")):
        bottom.set(qn(k), v)
    border.append(bottom)
    _insert_ordered(paragraph._p.get_or_add_pPr(), border, _PPR_ORDER)


def _bullet_items(children) -> list[str]:
    """The entries of a cell written as "• a • b" (plain text only), else []."""
    if any(tok.type not in ("text", "softbreak") for tok in children or []):
        return []
    text = "".join(tok.content if tok.type == "text" else " " for tok in children or []).strip()
    if not text.startswith("• "):
        return []
    return [i.strip() for i in re.split(r"(?:^|\s)•\s+", text) if i.strip()]


def _text_of(children) -> str:
    return "".join(tok.content for tok in children or [] if tok.type in ("text", "code_inline"))


_CHAR_IN = 0.085          # width of one character at the table's 8.5 pt, in inches (bold headers included)
_CELL_PAD_IN = 0.16       # cell margins


def _column_widths(rows: list[tuple[bool, list]], cols: int, page_width_in: float = PAGE_WIDTH_IN) -> list:
    """Column widths: each column at least as wide as its longest word (up to 16
    characters — longer ones, like rule ids, wrap at their hyphens), and the rest
    of the page shared in proportion to how much text the column holds. A short
    column ("Observed", "Medium") stays narrow without breaking its words."""
    minimum, weight = [], []
    for c in range(cols):
        texts = [_text_of(cells[c]) for _, cells in rows if c < len(cells)]
        longest_word = max((len(w) for t in texts for w in t.split()), default=4)
        average = sum(len(t) for t in texts) / max(len(texts), 1)
        # A short column keeps its words whole; a prose column wraps, so it needs less.
        cap = 16 if average <= 14 else 12
        minimum.append(min(longest_word, cap) * _CHAR_IN + _CELL_PAD_IN)
        weight.append(max(average - min(longest_word, cap), 1.0) ** 0.85)
    spare = page_width_in - sum(minimum)
    if spare <= 0:                                      # too many columns: shrink the minimums to fit
        return [Inches(page_width_in * m / sum(minimum)) for m in minimum]
    return [Inches(m + spare * w / sum(weight)) for m, w in zip(minimum, weight)]


class _Writer:
    def __init__(self, doc, images: dict[str, Path], page_width_in: float = PAGE_WIDTH_IN,
                 image_height_in: float = MAX_IMAGE_HEIGHT_IN):
        self.doc, self.images = doc, images
        self.page_width_in, self.image_height_in = page_width_in, image_height_in
        self.paragraph = None
        self.lists: list[dict] = []       # open lists: {"ordered", "number"}
        self.quote = 0
        self.item_paragraphs = 0          # paragraphs written in the current list item

    # ── inline ──────────────────────────────────────────────────────────────
    def inline(self, paragraph, children, size: Pt | None = None, bold: bool = False) -> None:
        state = {"bold": bold, "italic": False, "strike": False, "link": None}
        for tok in children or []:
            t = tok.type
            if t in ("strong_open", "strong_close"):
                state["bold"] = t == "strong_open" or bold
            elif t in ("em_open", "em_close"):
                state["italic"] = t == "em_open"
            elif t in ("s_open", "s_close"):
                state["strike"] = t == "s_open"
            elif t == "link_open":
                href = str(tok.attrs.get("href", ""))
                state["link"] = href if _LINKABLE.match(href) else None
            elif t == "link_close":
                state["link"] = None
            elif t == "softbreak":
                self._run(paragraph, " ", state, size)
            elif t == "hardbreak":
                paragraph.add_run().add_break()
            elif t == "code_inline":
                run = self._run(paragraph, tok.content, state, size)
                run.font.name = MONO
                run.font.color.rgb = RGBColor(0xA3, 0x1F, 0x34)
            elif t == "image":
                self.image(paragraph, str(tok.attrs.get("src", "")), tok.content or "")
            elif t in ("text", "html_inline"):
                self._run(paragraph, tok.content, state, size)

    def _run(self, paragraph, text: str, state: dict, size):
        text = _clean(text)
        if state["link"]:
            return self._hyperlink(paragraph, text, state["link"], size)
        run = paragraph.add_run(text)
        run.bold = state["bold"] or None
        run.italic = state["italic"] or None
        if state["strike"]:
            run.font.strike = True
        if size:
            run.font.size = size
        return run

    def _hyperlink(self, paragraph, text: str, url: str, size):
        rel = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rel)
        run = paragraph.add_run(text)
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0xA8)
        run.font.underline = True
        if size:
            run.font.size = size
        paragraph._p.remove(run._r)
        link.append(run._r)
        paragraph._p.append(link)
        return run

    def image(self, paragraph, src: str, alt: str) -> None:
        path = self.images.get(src)
        if path is None or not path.is_file():
            paragraph.add_run(f"[image: {_clean(alt) or src}]").italic = True
            return
        size = _png_size(path)
        width_in = self.page_width_in
        dpi = _png_dpi(path)
        if size and size[0] and dpi:                 # an image that declares its size (a diagram): no larger
            width_in = min(width_in, size[0] / dpi)
        width = Inches(width_in)
        if size and size[0]:
            height_in = width_in * size[1] / size[0]
            if height_in > self.image_height_in:
                width = Inches(width_in * self.image_height_in / height_in)
        paragraph.add_run().add_picture(str(path), width=width)

    # ── blocks ──────────────────────────────────────────────────────────────
    def new_paragraph(self):
        if self.lists:
            depth = len(self.lists)
            current = self.lists[-1]
            if self.item_paragraphs:                      # a further paragraph of the same item
                p = self.doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25 * depth + 0.25)
            elif current["ordered"]:
                p = self.doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25 * depth + 0.25)
                p.paragraph_format.first_line_indent = Inches(-0.25)
                p.add_run(f"{current['number']}. ")
            else:
                p = self.doc.add_paragraph(style="List Bullet" + ("" if depth == 1 else f" {min(depth, 3)}"))
            self.item_paragraphs += 1
        elif self.quote:
            p = self.doc.add_paragraph(style="Quote")
        else:
            p = self.doc.add_paragraph()
        return p

    def code(self, text: str) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.1)
        p.paragraph_format.space_after = Pt(6)
        _shade(p._p, "F3F5F9")
        lines = _clean(text).rstrip("\n").split("\n")
        for i, line in enumerate(lines):
            run = p.add_run(line)
            run.font.name = MONO
            run.font.size = Pt(8)
            if i < len(lines) - 1:
                run.add_break(WD_BREAK.LINE)

    def table(self, rows: list[tuple[bool, list]]) -> None:
        if not rows:
            return
        cols = max(len(cells) for _, cells in rows)
        table = self.doc.add_table(rows=len(rows), cols=cols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        widths = _column_widths(rows, cols, self.page_width_in)
        table.autofit = False
        for col, width in zip(table.columns, widths):
            col.width = width
        # Cells straight from the XML: python-docx's row.cells recomputes the
        # whole grid on every call, which is quadratic on a 2,000-row catalog.
        for tr, (header, cells) in zip(table._tbl.tr_lst, rows):
            # A row never splits across pages; the header row repeats on each page.
            tr_pr = tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:cantSplit"))
            if header:
                tr_pr.append(OxmlElement("w:tblHeader"))
            row_cells = [_Cell(tc, table) for tc in tr.tc_lst]
            for c in range(cols):
                cell = row_cells[c]
                cell.width = widths[c]
                p = cell.paragraphs[0]
                if c < len(cells):
                    items = _bullet_items(cells[c])
                    if items and not header:
                        # "• a • b": one paragraph per entry.
                        for k, item in enumerate(items):
                            para = p if k == 0 else cell.add_paragraph()
                            run = para.add_run(_clean(f"• {item}"))
                            run.font.size = Pt(8.5)
                    else:
                        self.inline(p, cells[c], size=Pt(8.5), bold=header)
                if header:
                    _shade(cell._tc, "E8EDF5")
        self.doc.add_paragraph()

    def render(self, markdown: str, top_level: int = 1, title: str = "") -> None:
        """`top_level`: the Word heading level the document's highest heading gets.
        `title`: the section's own heading — a first heading repeating it is left out."""
        tokens = _MD.parse(_COMMENT_LINE.sub("", markdown or ""))
        if title and len(tokens) > 2 and tokens[0].type == "heading_open" \
                and tokens[1].content.strip().lower() == title.strip().lower():
            tokens = tokens[3:]
        levels = [int(tok.tag[1]) for tok in tokens if tok.type == "heading_open"]
        heading_offset = top_level - min(levels) if levels else 0
        i = 0
        table_rows: list[tuple[bool, list]] | None = None
        cells: list = []
        in_head = False
        while i < len(tokens):
            tok = tokens[i]
            t = tok.type
            if t == "heading_open":
                level = max(1, min(9, int(tok.tag[1]) + heading_offset))
                p = self.doc.add_heading("", level=level)
                self.inline(p, tokens[i + 1].children)
                i += 3
                continue
            if t == "paragraph_open" and table_rows is None:
                inline = tokens[i + 1]
                kids = inline.children or []
                p = self.new_paragraph()
                if len(kids) == 1 and kids[0].type == "image":
                    p.paragraph_format.keep_with_next = False
                self.inline(p, kids)
                i += 3
                continue
            if t == "bullet_list_open":
                self.lists.append({"ordered": False, "number": 0})
            elif t == "ordered_list_open":
                self.lists.append({"ordered": True, "number": int(tok.attrs.get("start", 1) or 1) - 1})
            elif t in ("bullet_list_close", "ordered_list_close"):
                self.lists.pop()
                if not self.lists:
                    self.item_paragraphs = 0
            elif t == "list_item_open":
                if self.lists and self.lists[-1]["ordered"]:
                    self.lists[-1]["number"] = int(tok.info) if tok.info.isdigit() else self.lists[-1]["number"] + 1
                self.item_paragraphs = 0
            elif t == "list_item_close":
                self.item_paragraphs = 0
            elif t == "blockquote_open":
                self.quote += 1
            elif t == "blockquote_close":
                self.quote -= 1
            elif t in ("fence", "code_block"):
                self.code(tok.content)
            elif t == "hr":
                _rule(self.doc.add_paragraph())
            elif t == "html_block":
                if _clean(tok.content).strip():
                    self.code(tok.content)
            elif t == "table_open":
                table_rows = []
            elif t == "thead_open":
                in_head = True
            elif t == "thead_close":
                in_head = False
            elif t == "tr_open":
                cells = []
            elif t == "tr_close" and table_rows is not None:
                table_rows.append((in_head, cells))
            elif t == "inline" and table_rows is not None:
                cells.append(tok.children)
            elif t == "table_close":
                self.table(table_rows or [])
                table_rows = None
            i += 1


def _widest_table(markdown: str) -> int:
    """The most columns any table in `markdown` has (0 without tables)."""
    widest, cells, in_row = 0, 0, False
    for tok in _MD.parse(markdown or ""):
        if tok.type == "tr_open":
            cells, in_row = 0, True
        elif tok.type in ("th_open", "td_open") and in_row:
            cells += 1
        elif tok.type == "tr_close":
            widest, in_row = max(widest, cells), False
    return widest


def to_docx(title: str, sections: list[tuple[str | None, str]], images: dict[str, Path] | None = None,
            subtitle: str = "") -> bytes:
    """A Word document: `title`, then each (heading, markdown) section — a heading
    starts a new page and the section's own headings sit below it."""
    doc = Document()
    landscape = any(_widest_table(markdown) >= WIDE_TABLE_COLUMNS for _, markdown in sections)
    for s in doc.sections:
        if landscape:
            s.orientation = WD_ORIENT.LANDSCAPE
            s.page_width, s.page_height = s.page_height, s.page_width
        s.left_margin = s.right_margin = Inches(0.75)
        s.top_margin = s.bottom_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    doc.core_properties.title = _clean(title)
    doc.core_properties.author = "Stella Modernizer"
    doc.add_heading(_clean(title), level=0)
    if subtitle:
        doc.add_paragraph(_clean(subtitle)).runs[0].italic = True
    writer = _Writer(doc, images or {}, *((LANDSCAPE_WIDTH_IN, LANDSCAPE_IMAGE_HEIGHT_IN) if landscape
                                           else (PAGE_WIDTH_IN, MAX_IMAGE_HEIGHT_IN)))
    first = True
    for heading, markdown in sections:
        if heading:
            if not first:
                doc.add_page_break()
            doc.add_heading(_clean(heading), level=1)
        writer.render(markdown or "_Not produced for this run._", top_level=2 if heading else 1, title=heading or title)
        first = False
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
